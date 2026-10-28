#!/usr/bin/env python3
"""
Markdown cover letter -> plain-text DOCX (Calibri, no bold/italic).

Usage:
  python scripts/render-cover-letter-docx.py cover-letters/foo-cover-letter-2026-05-26.md
  python scripts/render-cover-letter-docx.py --all
  python scripts/render-cover-letter-docx.py --batch 2026-05-26-oregon-sor
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent.parent
COVER_DIR = ROOT / "cover-letters"
FONT_NAME = "Calibri"
BODY_SIZE = Pt(11)


def strip_md_inline(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def strip_trailing_footnote(md: str) -> str:
    return re.sub(r"\n---\n\n\*[\s\S]*$", "", md, flags=re.M).strip()


def split_header_body(md: str) -> tuple[str, str]:
    cleaned = strip_trailing_footnote(md)
    parts = cleaned.split("\n---\n", 1)
    if len(parts) == 2 and "Dear " not in parts[0]:
        return parts[0].strip(), parts[1].strip()
    return "", cleaned


def header_meta_lines(header_md: str) -> list[str]:
    lines: list[str] = []
    for raw in header_md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            lines.append(strip_md_inline(line[2:]))
            continue
        if line.startswith("**PDF:"):
            continue
        if line.startswith("**Posting:") or line.startswith("**JD:"):
            lines.append(strip_md_inline(line))
    return lines


def add_plain_paragraph(doc: Document, text: str, *, bullet: bool = False) -> None:
    style = "List Bullet" if bullet else None
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = BODY_SIZE
    run.bold = False
    run.italic = False


def body_to_docx(doc: Document, body_md: str) -> None:
    blocks = re.split(r"\n\n+", body_md.strip())
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if all(l.strip().startswith("- ") for l in lines if l.strip()):
            for line in lines:
                s = line.strip()
                if s.startswith("- "):
                    add_plain_paragraph(doc, strip_md_inline(s[2:]), bullet=True)
            continue
        if len(lines) == 1:
            add_plain_paragraph(doc, strip_md_inline(block.replace("\n", " ")))
            continue
        for line in lines:
            s = line.strip()
            if not s:
                continue
            if s.startswith("- "):
                add_plain_paragraph(doc, strip_md_inline(s[2:]), bullet=True)
            else:
                add_plain_paragraph(doc, strip_md_inline(s))


def render_one(md_path: Path) -> Path:
    md = md_path.read_text(encoding="utf-8")
    header_md, body_md = split_header_body(md)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE

    for line in header_meta_lines(header_md):
        add_plain_paragraph(doc, line)
    if header_meta_lines(header_md) and body_md:
        add_plain_paragraph(doc, "")

    body_to_docx(doc, body_md)

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def iter_cover_letters() -> list[Path]:
    return sorted(COVER_DIR.glob("*-cover-letter-*.md"))


def main(argv: list[str]) -> int:
    args = argv[1:]
    if not args:
        print(
            "Usage: python scripts/render-cover-letter-docx.py <file.md> | --all | --batch <slug>",
            file=sys.stderr,
        )
        return 1

    paths: list[Path] = []
    if args[0] == "--all":
        paths = iter_cover_letters()
    elif args[0] == "--batch" and len(args) > 1:
        slug = args[1]
        paths = sorted(COVER_DIR.glob(f"*-cover-letter-{slug}.md"))
        if not paths:
            paths = sorted(COVER_DIR.glob(f"*-{slug}*-cover-letter-*.md"))
        if not paths:
            paths = sorted(COVER_DIR.glob(f"*{slug}*cover-letter*.md"))
    else:
        p = Path(args[0])
        if not p.is_absolute():
            p = ROOT / p
        paths = [p]

    if not paths:
        print("No cover letter files matched.", file=sys.stderr)
        return 1

    for md_path in paths:
        if not md_path.exists():
            print(f"Skip (missing): {md_path}", file=sys.stderr)
            continue
        out = render_one(md_path)
        print(f"Wrote {out.relative_to(ROOT)}")

    print(f"\nDone: {len(paths)} DOCX file(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
